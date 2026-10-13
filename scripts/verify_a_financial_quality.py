"""Generate and verify an independent A candidate against existing official evidence.

Only cached filings and local source bytes are read. Database writes and network
access are prohibited. Existing company A/B/C files are never outputs.
"""
import argparse
from dataclasses import asdict
import hashlib
import json
from pathlib import Path
import socket
from unittest.mock import patch

import audit_fiscal_pipeline as audit


def digest_tree(root, tickers):
    paths = [root / 'database.json']
    paths.extend(p for p in (root / 'metadata').rglob('*') if p.is_file())
    for ticker in tickers:
        paths.extend(p for p in (root / 'filings' / ticker).rglob('*') if p.is_file())
    return {str(p.relative_to(root)): hashlib.sha256(p.read_bytes()).hexdigest() for p in paths}


def table_map(table):
    return {key: dict(zip(table['quarter_labels'], values))
            for key, values in zip(table['concepts'], table['values'])}


def verify_peer_changes(ticker, changes, result):
    """Every changed old cell must be nonadditive Q4 or an independently checked BS scope repair."""
    fg, fc = audit.fg, audit.fc
    bs_sources = []
    for path in fc.ticker_dir(ticker).glob('*.json'):
        if not fc.ACCESSION_RE.fullmatch(path.stem):
            continue
        entry = fc.load_filing(ticker, path.stem, fc.cached_cik(ticker))
        fin = fc.cached_filing(entry).financials
        stmt = fin.balance_sheet() if fin else None
        if stmt is None:
            continue
        df = stmt.to_dataframe()
        columns = [c for c in df.columns if fg._col_to_period_end(c)]
        if columns:
            bs_sources.append((path.stem, entry['form'], columns[0], df[fg._consolidated_mask(df)]))
    verified = []
    for change in changes:
        key, period, after = change['metric'], change['period'], change['after']
        if key in ('Basic EPS', 'Diluted EPS', 'Basic Shares', 'Diluted Shares'):
            assert period.endswith('Q4') and after is None, change
            verified.append({**change, 'reason': 'nonadditive Q4 safely missing; requires official standalone source'})
            continue
        assert key in ('Total Liabilities', 'Total Non-current Liabilities'), ('Unexpected peer regression', ticker, change)
        table = next(t for t in result['tables'] if t['sheet_name'] == change['sheet'])
        col = table['quarter_labels'].index(period)
        end = table['period_ends'][col]
        values = table_map(table)
        tl = values['Total Liabilities'][period]
        if key == 'Total Non-current Liabilities':
            assert after == tl - values['Total Current Liabilities'][period], change
        raw_checks = []
        for accession, source_form, column, rows in bs_sources:
            form = '10-K' if period.endswith('Q4') or 'Q' not in period else '10-Q'
            if source_form != form or column[:10] != end:
                continue
            def originals(concept):
                return set(v for v in rows.loc[rows['concept'].eq('us-gaap_' + concept), column].map(fg._to_python_val) if v is not None)
            reported = originals('Liabilities')
            if reported:
                assert reported == {tl}, (ticker, period, 'exact source liabilities', reported, tl)
                raw_checks.append(accession)
                continue
            total = originals('LiabilitiesAndStockholdersEquity')
            equity = originals('StockholdersEquityIncludingPortionAttributableToNoncontrollingInterest')
            if not equity and not rows['concept'].astype(str).str.contains('MinorityInterest|NoncontrollingInterest').any():
                equity = originals('StockholdersEquity')
            if len(total) == len(equity) == 1:
                assert next(iter(total)) - next(iter(equity)) == tl, (ticker, period, total, equity, tl)
                raw_checks.append(accession)
        assert raw_checks, (ticker, period, 'No current original BS source check')
        verified.append({**change, 'reason': 'original same-period liabilities/equity scope', 'accessions': raw_checks})
    return verified


def run(args):
    args.output.mkdir(parents=True, exist_ok=True)
    fc, fg = audit.fc, audit.fg
    from quarterly_source import read_quarter_source_pack
    from database import database_root
    from excel_writer import write_statements
    from output_tables import append_ratio_table
    from openpyxl import load_workbook

    root = database_root()
    before = digest_tree(root, [args.ticker, *args.peers])
    print('Source database hashes captured:', len(before), flush=True)
    source_records = read_quarter_source_pack(args.source_pack, fc.cached_cik(args.ticker))
    (args.output / 'official-source-records.json').write_text(json.dumps(source_records, indent=2), encoding='utf-8')
    def forbidden(*a, **kw):
        raise AssertionError('Network/database mutation forbidden during candidate validation')
    with patch.object(socket, 'create_connection', side_effect=forbidden), \
         patch.object(fc, 'save_filing', side_effect=forbidden):
        result = audit.audit(args.ticker, args.output, require_official_metadata=True,
                             metadata_dir=args.metadata_dir, quarter_source_pack=args.source_pack)
        tables = [fg.StatementTable(**t) for t in result['tables']]
        append_ratio_table(tables)
        candidate = args.output / f'{args.ticker} A candidate.xlsx'
        if candidate.exists():
            raise FileExistsError('Candidate path already exists; use a new output directory')
        write_statements(tables, candidate)
        # Peer runs intentionally have no supplement, verifying the shared logic.
        peer_results = {t: audit.audit(t, args.output, require_official_metadata=True,
                                     metadata_dir=args.metadata_dir) for t in args.peers}
    after = digest_tree(root, [args.ticker, *args.peers])
    assert before == after, 'Database bytes changed'

    quarterly = next(t for t in result['tables'] if t['sheet_name'] == 'Data_Financials(Q)')
    values = table_map(quarterly)
    evidence = json.loads(args.official_evidence.read_text(encoding='utf-8'))
    keys = {'revenue': 'Revenue', 'total_cost': 'Cost of Revenue', 'gross_profit': 'Gross Profit',
            'operating_income': 'Operating Income', 'pretax': 'Pre-tax Income', 'tax': 'Income Tax',
            'equity_income': 'Equity Method Income/(Loss)', 'net_income': 'Net Income',
            'basic_eps': 'Basic EPS', 'diluted_eps': 'Diluted EPS',
            'basic_shares': 'Basic Shares', 'diluted_shares': 'Diluted Shares',
            'discontinued': 'Discontinued Operations, net of tax'}
    checks, failures, absent = [], [], []
    wb = load_workbook(candidate, read_only=True, data_only=False)
    ws = wb['Data_Financials(Q)']
    row_map = {ws.cell(r, 1).value: r for r in range(1, ws.max_row + 1)}
    for record in evidence['selected']:
        period = record['period']
        col = quarterly['quarter_labels'].index(period) + 4
        source_path = args.source_pack / record['file']
        assert hashlib.sha256(source_path.read_bytes()).hexdigest() == record['sha256']
        original = source_records[record['period_end']]
        # A supported later comparative table can carry the same standalone
        # values when the original release has an unsupported HTML layout.
        # Verify its bytes through the reader, and require numeric agreement
        # with the independently selected original release below.
        for raw_key, key in keys.items():
            expected = record['values'].get(raw_key)
            if expected is None:
                absent.append(dict(period=period, metric=key, reason='not separately presented'))
                continue
            expected *= 1 if key.endswith('EPS') else 1_000_000
            # Independently selected handoff numbers must also agree with raw parser.
            if key in original['values']:
                assert abs(original['values'][key] - expected) < 1e-5, (period, key, 'raw parser disagreement')
            actual = values.get(key, {}).get(period)
            workbook = ws.cell(row_map[key], col).value
            workbook_expected = expected if key.endswith('EPS') else expected / 1_000_000
            ok = actual is not None and abs(actual - expected) < 1e-5 and workbook is not None and abs(workbook - workbook_expected) < 1e-8
            check = dict(period=period, metric=key, expected=expected, actual=actual, workbook=workbook, passed=ok,
                         original_accession=record['accession'], parsed_accession=original['accession'])
            checks.append(check)
            if not ok:
                failures.append(check)
    wb.close()
    changes, peer_checks = {}, {}
    for ticker, current in peer_results.items():
        baseline = json.loads((args.baseline / (ticker + '.json')).read_text(encoding='utf-8'))
        changed = []
        for table in current['tables']:
            if table['sheet_name'] not in ('Data_Financials(Q)', 'Data_Financials(Y)'):
                continue
            previous = next(t for t in baseline['tables'] if t['sheet_name'] == table['sheet_name'])
            old = table_map(previous)
            new = table_map(table)
            assert table['quarter_labels'] == previous['quarter_labels'], ticker + ' periods changed'
            for key in old:
                if key not in new:
                    raise AssertionError((ticker, key, 'row lost'))
                for period, old_value in old[key].items():
                    new_value = new[key][period]
                    if old_value != new_value:
                        changed.append(dict(sheet=table['sheet_name'], metric=key, period=period,
                                            before=old_value, after=new_value))
        changes[ticker] = changed
        peer_checks[ticker] = verify_peer_changes(ticker, changed, current)
    report = dict(status='PASS' if not failures else 'FAIL', checks=len(checks),
                  official_periods=len(evidence['selected']), failures=failures,
                  not_presented=absent, gaps=result['gaps'], peer_changes=changes,
                  database_unchanged=True, database_hash_scope=[args.ticker, *args.peers, 'metadata', 'database.json'],
                  database_files=len(before), candidate=str(candidate.resolve()),
                  candidate_sha256=hashlib.sha256(candidate.read_bytes()).hexdigest())
    (args.output / 'verification.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
    (args.output / 'official-checks.json').write_text(json.dumps(checks, indent=2), encoding='utf-8')
    (args.output / 'peer-checks.json').write_text(json.dumps(peer_checks, indent=2), encoding='utf-8')
    print(json.dumps({k: v for k, v in report.items() if k not in ('not_presented', 'peer_changes', 'gaps', 'failures')}, indent=2))
    print('failures', len(failures), 'peer changed cells', {k: len(v) for k, v in changes.items()})
    return 1 if failures else 0


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--ticker', default='AMD')
    parser.add_argument('--source-pack', type=Path, required=True)
    parser.add_argument('--official-evidence', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--baseline', type=Path, required=True)
    parser.add_argument('--metadata-dir', type=Path, required=True)
    parser.add_argument('--peers', nargs='*', default=['AAPL', 'NVDA', 'MSFT', 'MRVL'])
    raise SystemExit(run(parser.parse_args()))
