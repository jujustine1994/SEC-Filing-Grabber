"""Verify saved candidate values after native Excel recalculation and update its hash."""
import argparse
import hashlib
import json
from pathlib import Path
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('folder', type=Path)
parser.add_argument('--summary', type=Path)
args = parser.parse_args()
folder = args.folder
report = json.loads((folder / 'verification.json').read_text(encoding='utf-8'))
checks = json.loads((folder / 'official-checks.json').read_text(encoding='utf-8'))
pipeline = json.loads((folder / 'AMD.json').read_text(encoding='utf-8'))
native = json.loads((folder / 'native-excel.json').read_text(encoding='utf-8-sig'))
assert native['status'] == report['status'] == 'PASS'
quarterly = next(t for t in pipeline['tables'] if t['sheet_name'] == 'Data_Financials(Q)')
path = Path(report['candidate'])
wb = load_workbook(path, read_only=True, data_only=True)
ws = wb['Data_Financials(Q)']
rows = {}
for row in ws.iter_rows(min_col=1, max_col=1):
    rows.setdefault(row[0].value, row[0].row)
for check in checks:
    col = quarterly['quarter_labels'].index(check['period']) + 4
    actual = ws.cell(rows[check['metric']], col).value
    assert actual is not None and abs(actual - check['workbook']) < 1e-8, check
wb.close()
report.setdefault('pre_excel_candidate_sha256', report['candidate_sha256'])
report['candidate_sha256'] = hashlib.sha256(path.read_bytes()).hexdigest()
report['native_excel'] = native
report['post_excel_checks'] = len(checks)
(folder / 'verification.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
if args.summary:
    junit = folder.parent / 'final-suite.xml'
    suite = ET.parse(junit).getroot().find('testsuite')
    assert suite is not None and suite.get('failures') == suite.get('errors') == '0'
    summary = {k: report[k] for k in ('status', 'checks', 'official_periods', 'database_unchanged',
        'database_hash_scope', 'database_files', 'candidate', 'candidate_sha256',
        'pre_excel_candidate_sha256', 'native_excel', 'post_excel_checks', 'gaps')}
    summary['offline_tests'] = {'passed': int(suite.get('tests')), 'failures': 0,
                               'junit_sha256': hashlib.sha256(junit.read_bytes()).hexdigest()}
    summary['peer_changed_cells'] = {ticker: len(rows) for ticker, rows in report['peer_changes'].items()}
    summary['comparative_source_periods'] = sorted({c['period'] for c in checks
        if c['original_accession'] != c['parsed_accession']})
    summary['official_source_checks'] = str((folder / 'official-checks.json').resolve())
    summary['peer_source_checks'] = str((folder / 'peer-checks.json').resolve())
    summary['basis'] = 'Existing cached-input pipeline; official original releases independently check parsed standalone tables; sampled peers, not an all-company official audit'
    args.summary.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(dict(status='PASS', post_excel_checks=len(checks), candidate_sha256=report['candidate_sha256'])))
